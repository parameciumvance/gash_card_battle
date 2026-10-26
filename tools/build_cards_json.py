"""從 data/cards_ja.csv(日文權威來源)轉換產出 data/cards.json。

- `power` 單一字串欄位解析回 {base/bonus/special/per_heads} 結構。
- `attr_name` 只在術卡填入 `attr_ja`,魔物/搭檔卡一律 None(僅 M-023 讀取,不承擔顯示職責)。
- `related_mamodo`:術/搭檔/事件卡直接取 `related_mamodo_ja`;魔物卡以 `name_ja` 去除
  括號後綴(如「（変身後）」)推導,使變身/形態卡與基礎家族名一致。
- `effect_icon`:バトル/非バトル/ジャマー 轉為 battle/nonbattle/jammer,空字串為 None。
- `sets` 沿用轉換前 `data/cards.json` 同卡號的舊值;無舊值的新卡取 xlsx 的 Sets 欄,兩者都沒有則為 []。
- 不含卡圖連結:卡圖來源是 `tools/download_images.py` 的資料夾。

用法: python tools/build_cards_json.py
"""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CSV_PATH = ROOT / "data/cards_ja.csv"
XLSX = ROOT / "openspec/specs/card-data/Zatch Bell CCG List for TTS.xlsx"
OUT = ROOT / "data/cards.json"

EXPECTED_COUNT = 237  # 第一彈+Level 2 共 135(含 S-042)+ 第三彈 102

FORM_SUFFIX_RE = re.compile(r"（[^）]*）")
POWER_X_RE = re.compile(r"^(\d+)×$")
POWER_BONUS_RE = re.compile(r"^\+(\d+)$")
POWER_BASE_RE = re.compile(r"^(\d+)$")

EFFECT_ICON_MAP = {"バトル": "battle", "非バトル": "nonbattle", "ジャマー": "jammer"}


def _none_if_blank(s: str) -> str | None:
    s = s.strip()
    return s or None


def _int_or_none(s: str) -> int | None:
    s = s.strip()
    return int(s) if s else None


def parse_power(value: str) -> dict:
    s = value.strip()
    if not s:
        return {}
    if s == "特殊":
        return {"special": True}
    m = POWER_X_RE.match(s)
    if m:
        return {"special": True, "per_heads": int(m.group(1))}
    m = POWER_BONUS_RE.match(s)
    if m:
        return {"bonus": int(m.group(1))}
    m = POWER_BASE_RE.match(s)
    if m:
        return {"base": int(m.group(1))}
    raise ValueError(f"無法解析 power 格式: {value!r}")


def derive_related_mamodo(row: dict) -> str | None:
    if row["type"] == "mamodo":
        base = FORM_SUFFIX_RE.sub("", row["name_ja"]).strip()
        return base or None
    return _none_if_blank(row["related_mamodo_ja"])


def derive_attr_name(row: dict) -> str | None:
    if row["type"] == "spell":
        return _none_if_blank(row["attr_ja"])
    return None


def derive_effect_icon(row: dict) -> str | None:
    icon = row["effect_icon_ja"].strip()
    if not icon:
        return None
    return EFFECT_ICON_MAP[icon]


def card_number(cell) -> str:
    """A 欄的卡號:HYPERLINK 公式(=HYPERLINK("url","E-001"))取顯示文字,其他取儲存格值。"""
    value = str(cell.value or "")
    if value.startswith("=HYPERLINK"):
        quoted = re.findall(r'"([^"]*)"', value)
        if len(quoted) >= 2:
            return quoted[-1]
        raise ValueError(f"無法解析卡號欄: {value!r}")
    return value.strip()


def load_tts(xlsx_path: Path) -> dict[str, dict]:
    """從 xlsx「The Table」工作表逐列讀取 卡號 → {sets},e/j 雙版本時 j 優先。
    sets 為 Sets 欄以「, 」分隔的產品(新卡沒有舊值時使用)。"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    ws = wb["The Table"]
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    sets_col = header.index("Sets")
    out: dict[str, dict] = {}
    versions: dict[str, str] = {}  # 基礎卡號 → 已入庫版本尾碼('' / 'e' / 'j')
    for cells in ws.iter_rows(min_row=2):
        raw_number = card_number(cells[0])
        m = re.fullmatch(r"([A-Z]+-\d+)([ej]?)", raw_number)
        if not m:
            continue
        number, suffix = m.group(1), m.group(2)
        if number in versions and not (suffix == "j" and versions[number] == "e"):
            continue
        versions[number] = suffix
        raw_sets = str(cells[sets_col].value or "")
        out[number] = {"sets": [x.strip() for x in raw_sets.split(", ") if x.strip()]}
    return out


def load_old_cards(path: Path) -> dict[str, dict]:
    if not path.exists():
        return {}
    raws = json.loads(path.read_text(encoding="utf-8"))
    return {c["number"]: c for c in raws}


def convert(csv_path: Path, tts: dict[str, dict], old_cards: dict[str, dict],
            expected: int | None = EXPECTED_COUNT) -> list[dict]:
    with csv_path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))

    cards = []
    for row in rows:
        old = old_cards.get(row["number"], {})
        info = tts.get(row["number"], {})
        card = (
            {
                "number": row["number"],
                "type": row["type"],
                "name_ja": row["name_ja"],
                "related_mamodo": derive_related_mamodo(row),
                "cost": _int_or_none(row["cost"]),
                "ad": _none_if_blank(row["ad"]),
                "class": row["class"],
                "attr_name": derive_attr_name(row),
                "effect_ja": row["effect_ja"],
                "effect_icon": derive_effect_icon(row),
                "power": parse_power(row["power"]),
                "damage": None if row["damage"] == "?" else _int_or_none(row["damage"]),
                # 沿用舊值;新卡取 TTS 卡表的 Sets 欄
                "sets": old["sets"] if "sets" in old else info.get("sets", []),
            }
        )
        if row["damage"] == "?":                 # 傷害依效果決定(卡面「ダメージ?」)
            card["damage_variable"] = True
        cards.append(card)

    cards.sort(key=lambda c: (c["number"][0], c["number"]))
    if expected is not None and len(cards) != expected:
        raise SystemExit(f"轉換出 {len(cards)} 種,預期 {expected} 種")
    return cards


def main() -> None:
    old_cards = load_old_cards(OUT)
    tts = load_tts(XLSX)
    cards = convert(CSV_PATH, tts, old_cards)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(cards, ensure_ascii=False, indent=1), encoding="utf-8")

    from collections import Counter

    counts = Counter(c["type"] for c in cards)
    print(f"已寫入 {OUT}: {len(cards)} 種 {dict(counts)}")


if __name__ == "__main__":
    main()
